"""Budget codes with no numeric actuals stay in the master and training snapshot."""

from __future__ import annotations

import tempfile
import unittest
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app.services.dataset_master import (
    master_identity_from_workbook_path,
    prepare_managed_master,
    resolve_training_snapshot,
)
from app.services.dataset_service import confirm_manual_edits
from app.services.dataset_workflow import analyze_upload, next_master_state


def _workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Historical Actuals"
    worksheet.append(
        ["ACT CODE", "ACT NAME", "Category", "January23 Actuals", "February23 Actuals", "March23 Actuals"]
    )
    worksheet.append([None, "Staff Cost", "Staff Cost", None, None, None])
    worksheet.append(["511101", "Salaries", "Staff Cost", 12, None, None])
    worksheet.append(["511199", "Unmeasured", "Staff Cost", None, None, None])
    workbook.save(path)
    workbook.close()


def _row_map(path: Path) -> tuple[list, dict[str, tuple]]:
    saved = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(saved.active.iter_rows(values_only=True))
    finally:
        saved.close()
    headers = list(rows[0])
    return headers, {str(row[0]): row for row in rows if row[0]}


class MasterCodePreservationTests(unittest.TestCase):
    def test_first_initialization_keeps_codes_with_no_actuals_and_the_month_range(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "historical_actuals_master.xlsx"
            _workbook(path)
            state = master_identity_from_workbook_path(path)
        codes = {item["budget_code"]: item for item in state["budget_codes"]}
        self.assertEqual(set(codes), {"511101", "511199"})
        self.assertEqual(codes["511199"]["description"], "Unmeasured")
        self.assertEqual(codes["511199"]["category"], "Staff Cost")
        self.assertEqual(state["earliest_month"], "2023-01")
        self.assertEqual(state["latest_month"], "2023-03")
        self.assertEqual(
            [(item["budget_code"], item["month"], item["amount"]) for item in state["records"]],
            [("511101", "2023-01", 12.0)],
        )

    def test_clearing_the_last_actual_keeps_the_code_and_month_range(self):
        state = {
            "records": [
                {
                    "budget_code": "511101",
                    "month": "2023-01",
                    "amount": 12.0,
                    "description": "Salaries",
                    "category": "Staff Cost",
                }
            ],
            "budget_codes": [
                {"budget_code": "511101", "description": "Salaries", "category": "Staff Cost"},
                {"budget_code": "511199", "description": "Unmeasured", "category": "Staff Cost"},
            ],
            "earliest_month": "2023-01",
            "latest_month": "2023-03",
        }
        cleared = next_master_state(
            state,
            [
                {
                    "operation": "CLEAR",
                    "budget_code": "511101",
                    "month": "2023-01",
                    "new_amount": None,
                    "new_description": "Salaries",
                    "new_category": "Staff Cost",
                }
            ],
        )
        self.assertEqual(cleared["records"], [])
        self.assertEqual(
            [item["budget_code"] for item in cleared["budget_codes"]],
            ["511101", "511199"],
        )
        self.assertEqual(cleared["earliest_month"], "2023-01")
        self.assertEqual(cleared["latest_month"], "2023-03")

    def test_export_writes_missing_actuals_instead_of_zeros(self):
        state = {
            "records": [
                {
                    "budget_code": "511101",
                    "month": "2023-01",
                    "amount": 0.0,
                    "description": "Salaries",
                    "category": "Staff Cost",
                }
            ],
            "budget_codes": [
                {"budget_code": "511101", "description": "Salaries", "category": "Staff Cost"},
                {"budget_code": "511199", "description": "Unmeasured", "category": "Staff Cost"},
            ],
            "earliest_month": "2023-01",
            "latest_month": "2023-03",
        }
        with tempfile.TemporaryDirectory() as tmp:
            with patch("app.services.dataset_storage.DATASET_STORAGE_DIR", Path(tmp)):
                exported = prepare_managed_master(
                    state["records"],
                    budget_codes=state["budget_codes"],
                    earliest_month=state["earliest_month"],
                    latest_month=state["latest_month"],
                )
                headers, rows = _row_map(exported)
        self.assertIn("February23 Actuals", headers)
        february = headers.index("February23 Actuals")
        january = headers.index("January23 Actuals")
        self.assertEqual(rows["511101"][january], 0)
        self.assertIsNone(rows["511101"][february])
        self.assertTrue(all(value is None for value in rows["511199"][3:]))
        self.assertNotIn(0, rows["511199"][3:])

    def test_retraining_snapshot_includes_the_missing_code(self):
        records = [
            {
                "budget_code": "511101",
                "month": "2023-01",
                "amount": 12.0,
                "description": "Salaries",
                "category": "Staff Cost",
            }
        ]
        identity = {
            "budget_codes": [
                {"budget_code": "511101", "description": "Salaries", "category": "Staff Cost"},
                {"budget_code": "511199", "description": "Unmeasured", "category": "Staff Cost"},
            ],
            "earliest_month": "2023-01",
            "latest_month": "2023-03",
        }
        with tempfile.TemporaryDirectory() as tmp:
            with (
                patch("app.services.dataset_storage.DATASET_STORAGE_DIR", Path(tmp)),
                patch("app.services.dataset_master.list_historical_actuals", return_value=records),
                patch("app.services.dataset_master.master_version", return_value=2),
                patch("app.services.dataset_master.load_master_identity", return_value=identity),
                patch("app.services.dataset_master.authoritative_workbook_path") as workbook_path,
            ):
                snapshot = resolve_training_snapshot()
            workbook_path.assert_not_called()
            headers, rows = _row_map(snapshot["path"])
        self.assertEqual(snapshot["source"], "database")
        self.assertEqual(snapshot["budget_code_count"], 2)
        self.assertEqual(snapshot["earliest_month"], "2023-01")
        self.assertEqual(snapshot["latest_month"], "2023-03")
        self.assertIn("511199", rows)
        self.assertTrue(all(value is None for value in rows["511199"][3:]))
        self.assertIn("February23 Actuals", headers)

    def test_manual_clear_confirmation_persists_the_code(self):
        state = {
            "records": [
                {
                    "budget_code": "511101",
                    "month": "2023-01",
                    "amount": 12.0,
                    "description": "Salaries",
                    "category": "Staff Cost",
                }
            ],
            "budget_codes": [
                {"budget_code": "511101", "description": "Salaries", "category": "Staff Cost"}
            ],
            "earliest_month": "2023-01",
            "latest_month": "2023-03",
            "version": 3,
            "source": "database",
        }
        with (
            patch("app.services.dataset_service.mutation_scope", return_value=nullcontext()),
            patch("app.services.dataset_service.ensure_forecast_tables"),
            patch("app.services.dataset_service.load_master_state", return_value=state),
            patch("app.services.dataset_service.prepare_managed_master", return_value=Path("prepared.xlsx")),
            patch("app.services.dataset_service.publish_managed_master"),
            patch("app.services.dataset_service._invalidate_overview"),
            patch(
                "app.services.dataset_service.commit_master_state",
                return_value={"audit_id": 9, "version": 4},
            ) as commit,
        ):
            confirm_manual_edits(
                {
                    "name": "User",
                    "email": "user@example.com",
                    "master_version": 3,
                    "edits": [
                        {
                            "budget_code": "511101",
                            "month": "2023-01",
                            "amount": "-",
                            "expected_amount": 12,
                        }
                    ],
                }
            )
        saved = commit.call_args.kwargs
        self.assertEqual(saved["next_records"], [])
        self.assertEqual(saved["budget_codes"][0]["budget_code"], "511101")
        self.assertEqual(saved["earliest_month"], "2023-01")
        self.assertEqual(saved["latest_month"], "2023-03")

    def test_preserved_code_is_not_treated_as_new(self):
        analysis = analyze_upload(
            [
                {
                    "budget_code": "511199",
                    "description": "Unmeasured",
                    "category": "Staff Cost",
                    "stage_id": "row-1",
                    "amounts": {"2023-02": 4},
                }
            ],
            [],
            budget_codes=[{"budget_code": "511199", "description": "Unmeasured", "category": "Staff Cost"}],
            earliest_month="2023-01",
            latest_month="2023-03",
        )
        self.assertEqual(analysis["summary"]["new_budget_codes"], 0)
        self.assertEqual(analysis["summary"]["existing_codes"], 1)
        self.assertEqual(analysis["changes"][0]["operation"], "ADD")
        self.assertEqual(analysis["summary"]["latest_observed_month_after_save"], "2023-03")


if __name__ == "__main__":
    unittest.main()
