"""Focused Excel export-format tests. Forecast values are reused, not recalculated."""

from __future__ import annotations

import hashlib
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from app.routers import forecasting
from app.services.report_builder import (
    build_excel_forecast_model,
    build_forecast_report,
    excel_filename_stem,
    excel_report_title,
    month_column_headers,
)
from app.services.report_excel import render_forecast_excel
from app.tests.test_forecast_report import completed_payload

ENGINE_OUTPUT = Path(__file__).resolve().parents[2] / "model_training_engine" / "output"
PROTECTED_PICKLES = (
    ENGINE_OUTPUT / "all_budget_code_models.pkl",
    ENGINE_OUTPUT / "international_settlement_top12_models.pkl",
)


def _sha256(path: Path) -> str | None:
    if not path.exists():
        return None
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _month_range(start: str, end: str) -> list[str]:
    year, month = int(start[:4]), int(start[5:7])
    end_year, end_month = int(end[:4]), int(end[5:7])
    months = []
    while (year, month) <= (end_year, end_month):
        months.append(f"{year:04d}-{month:02d}")
        if month == 12:
            year += 1
            month = 1
        else:
            month += 1
    return months


def _code_rows(code: str, name: str, months: list[str], values: list[float | None]) -> list[dict]:
    rows = []
    for month, amount in zip(months, values):
        if amount is None:
            continue
        rows.append({
            "budget_code": code,
            "account_name": name,
            "month": month,
            "forecast_amount": amount,
        })
    return rows


def export_payload(**overrides):
    months = _month_range("2027-01", "2027-12")
    a_values = [10.0, 20.0, 15.0, 12.0, 8.0, 9.0, 11.0, 13.0, 14.0, 16.0, 17.0, 15.0]
    b_values = [30.0, 40.0, 25.0, 22.0, 18.0, 19.0, 21.0, 23.0, 24.0, 26.0, 27.0, 25.0]
    z_values = [0.0] * 12
    payload = completed_payload(
        category="All Categories",
        requested_start_month="2027-01",
        requested_end_month="2027-12",
        forecast_start="2027-01",
        forecast_end="2027-12",
        forecast_month_count=12,
        selected_account_count=5,
        selected_accounts=["521101", "521102", "Z01", "U01", "511101"],
        monthly_forecasts=[
            {"month": month, "forecast_amount": a_values[index] + b_values[index]}
            for index, month in enumerate(months)
        ],
        account_monthly_forecasts=(
            _code_rows("521101", "Repair A", months, a_values)
            + _code_rows("521102", "Repair B", months, b_values)
            + _code_rows("Z01", "Zero Policy Code", months, z_values)
            + _code_rows("511101", "IS One", months, [-2.5] + [1.0] * 11)
        ),
        budget_code_forecasts=[
            {
                "budget_code": "521101",
                "account_name": "Repair A",
                "category": "R&M-AMCs",
                "available": True,
                "forecast": dict(zip(months, a_values)),
                "forecast_type": "FITTED_MODEL",
                "production_status": "FITTED_MODEL",
                "display_status": "Available",
            },
            {
                "budget_code": "521102",
                "account_name": "Repair B",
                "category": "R&M-AMCs",
                "available": True,
                "forecast": dict(zip(months, b_values)),
                "forecast_type": "FITTED_MODEL",
                "production_status": "FITTED_MODEL",
                "display_status": "Available",
            },
            {
                "budget_code": "Z01",
                "account_name": "Zero Policy Code",
                "category": "R&M-AMCs",
                "available": True,
                "forecast": dict(zip(months, z_values)),
                "forecast_type": "ZERO_POLICY",
                "production_status": "ZERO_POLICY",
                "display_status": "Zero Policy",
            },
            {
                "budget_code": "U01",
                "account_name": "Unavailable Code",
                "category": "R&M-AMCs",
                "available": False,
                "forecast": None,
                "forecast_type": None,
                "production_status": "NO_VALID_WINNER",
                "display_status": "Forecast Unavailable",
            },
            {
                "budget_code": "511101",
                "account_name": "IS One",
                "category": "Int'l Settlement",
                "available": True,
                "forecast": dict(zip(months, [-2.5] + [1.0] * 11)),
                "forecast_type": "FITTED_MODEL",
                "production_status": "FITTED_MODEL",
                "display_status": "Available",
            },
        ],
    )
    payload.update(overrides)
    return payload


def _workbook(payload):
    report = build_forecast_report(payload)
    workbook = load_workbook(BytesIO(render_forecast_excel(report)))
    return workbook, report


def _sheet(payload):
    workbook, report = _workbook(payload)
    assert workbook.sheetnames == ["Budget Forecast"]
    return workbook["Budget Forecast"], report


def _row_for(sheet, value):
    for row in sheet.iter_rows(min_row=1, max_col=1):
        if row[0].value == value:
            return row[0].row
    return None


class ExcelForecastExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pickle_hashes = {str(path): _sha256(path) for path in PROTECTED_PICKLES}

    def test_workbook_contains_exactly_one_budget_forecast_sheet(self):
        workbook, _ = _workbook(export_payload())
        self.assertEqual(workbook.sheetnames, ["Budget Forecast"])
        self.assertEqual(len(workbook.sheetnames), 1)
        for name in (
            "Forecast Summary",
            "Monthly Forecast",
            "Budget Code Summary",
            "Account Monthly",
            "Yearly Forecast",
            "Historical Comparison",
        ):
            self.assertNotIn(name, workbook.sheetnames)

    def test_jan_dec_export_creates_twelve_month_columns(self):
        model = build_excel_forecast_model(export_payload())
        self.assertEqual(len(model["months"]), 12)
        self.assertEqual(model["months"][0], "2027-01")
        self.assertEqual(model["months"][-1], "2027-12")
        self.assertEqual(model["month_headers"], [
            "Jan", "Feb", "Mar", "Apr", "May", "Jun",
            "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
        ])
        sheet, _ = _sheet(export_payload())
        headers = [sheet.cell(4, column).value for column in range(1, 16)]
        self.assertEqual(headers[0], "ACT CODE")
        self.assertEqual(headers[1], "ACT NAME")
        self.assertEqual(headers[-1], "YTD")
        self.assertEqual(headers[2:14], model["month_headers"])

    def test_partial_period_export_creates_only_selected_month_columns(self):
        months = _month_range("2027-07", "2027-12")
        payload = export_payload(
            requested_start_month="2027-07",
            requested_end_month="2027-12",
            forecast_month_count=6,
            monthly_forecasts=[{"month": month, "forecast_amount": 1.0} for month in months],
            account_monthly_forecasts=_code_rows("521101", "Repair A", months, [10.0] * 6),
            selected_accounts=["521101"],
            selected_account_count=1,
            category="R&M-AMCs",
            budget_code_forecasts=[{
                "budget_code": "521101",
                "account_name": "Repair A",
                "category": "R&M-AMCs",
                "available": True,
                "forecast": dict(zip(months, [10.0] * 6)),
                "production_status": "FITTED_MODEL",
            }],
        )
        model = build_excel_forecast_model(payload)
        self.assertEqual(model["month_headers"], ["Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])
        self.assertEqual(len(model["months"]), 6)

    def test_cross_year_period_creates_correct_ordered_columns(self):
        months = _month_range("2029-11", "2030-02")
        payload = export_payload(
            requested_start_month="2029-11",
            requested_end_month="2030-02",
            forecast_month_count=4,
            monthly_forecasts=[{"month": month, "forecast_amount": 1.0} for month in months],
            account_monthly_forecasts=_code_rows("521101", "Repair A", months, [1.0, 2.0, 3.0, 4.0]),
            selected_accounts=["521101"],
            selected_account_count=1,
            category="R&M-AMCs",
            budget_code_forecasts=[{
                "budget_code": "521101",
                "account_name": "Repair A",
                "category": "R&M-AMCs",
                "available": True,
                "forecast": dict(zip(months, [1.0, 2.0, 3.0, 4.0])),
                "production_status": "FITTED_MODEL",
            }],
        )
        model = build_excel_forecast_model(payload)
        self.assertEqual(model["months"], ["2029-11", "2029-12", "2030-01", "2030-02"])
        self.assertEqual(month_column_headers(model["months"]), ["Nov 2029", "Dec 2029", "Jan 2030", "Feb 2030"])

    def test_dynamic_title_contains_exact_start_and_end_months(self):
        self.assertEqual(
            excel_report_title("2026-07", "2026-12"),
            "AI-Based Budget Forecast (July 2026 - December 2026)",
        )
        self.assertEqual(
            excel_report_title("2027-01", "2027-12"),
            "AI-Based Budget Forecast (January 2027 - December 2027)",
        )
        self.assertEqual(
            excel_report_title("2028-03", "2028-08"),
            "AI-Based Budget Forecast (March 2028 - August 2028)",
        )
        sheet, report = _sheet(export_payload())
        self.assertEqual(sheet["A1"].value, "AI-Based Budget Forecast (January 2027 - December 2027)")
        self.assertTrue(sheet["A1"].font.bold)
        self.assertEqual(report["excel_forecast"]["title"], sheet["A1"].value)

    def test_act_code_and_name_are_present(self):
        sheet, _ = _sheet(export_payload())
        row = _row_for(sheet, "521101")
        self.assertIsNotNone(row)
        self.assertEqual(sheet.cell(row, 1).value, "521101")
        self.assertEqual(sheet.cell(row, 2).value, "Repair A")
        self.assertEqual(sheet.cell(row, 1).number_format, "@")

    def test_budget_codes_appear_under_the_correct_category(self):
        model = build_excel_forecast_model(export_payload())
        by_category = {section["category"]: [row["budget_code"] for row in section["rows"]] for section in model["sections"]}
        self.assertEqual(by_category["R&M-AMCs"], ["521101", "521102", "Z01", "U01"])
        self.assertEqual(by_category["Int'l Settlement"], ["511101"])
        sheet, _ = _sheet(export_payload())
        rm_row = _row_for(sheet, "R&M-AMCs")
        is_row = _row_for(sheet, "Int'l Settlement")
        code_row = _row_for(sheet, "521101")
        other_row = _row_for(sheet, "511101")
        self.assertLess(rm_row, code_row)
        self.assertLess(code_row, is_row)
        self.assertLess(is_row, other_row)

    def test_budget_code_ytd_equals_sum_of_exported_months(self):
        model = build_excel_forecast_model(export_payload())
        row = next(item for item in model["sections"][0]["rows"] if item["budget_code"] == "521101")
        self.assertEqual(row["ytd"], sum(row["months"]))
        self.assertEqual(row["ytd"], 160.0)
        sheet, _ = _sheet(export_payload())
        excel_row = _row_for(sheet, "521101")
        self.assertEqual(sheet.cell(excel_row, 15).value, 160.0)

    def test_category_january_total_equals_available_january_forecasts(self):
        model = build_excel_forecast_model(export_payload())
        section = next(item for item in model["sections"] if item["category"] == "R&M-AMCs")
        january = [row["months"][0] for row in section["rows"] if row["months"][0] is not None]
        self.assertEqual(january, [10.0, 30.0, 0.0])
        self.assertEqual(section["totals"][0], 40.0)
        sheet, _ = _sheet(export_payload())
        total_row = _row_for(sheet, "TOTAL")
        self.assertEqual(sheet.cell(total_row, 3).value, 40.0)

    def test_category_february_total_equals_available_february_forecasts(self):
        model = build_excel_forecast_model(export_payload())
        section = next(item for item in model["sections"] if item["category"] == "R&M-AMCs")
        self.assertEqual(section["totals"][1], 60.0)
        sheet, _ = _sheet(export_payload())
        total_row = _row_for(sheet, "TOTAL")
        self.assertEqual(sheet.cell(total_row, 4).value, 60.0)

    def test_category_total_row_ytd_is_blank(self):
        model = build_excel_forecast_model(export_payload())
        for section in model["sections"]:
            self.assertIsNone(section["total_ytd"])
        sheet, _ = _sheet(export_payload())
        total_row = _row_for(sheet, "TOTAL")
        self.assertIsNone(sheet.cell(total_row, 15).value)

    def test_zero_policy_remains_visible_with_zero_values(self):
        model = build_excel_forecast_model(export_payload())
        row = next(item for item in model["sections"][0]["rows"] if item["budget_code"] == "Z01")
        self.assertTrue(row["zero_policy"])
        self.assertEqual(row["months"], [0.0] * 12)
        self.assertEqual(row["ytd"], 0.0)
        sheet, _ = _sheet(export_payload())
        excel_row = _row_for(sheet, "Z01")
        self.assertEqual(sheet.cell(excel_row, 3).value, 0.0)
        self.assertEqual(sheet.cell(excel_row, 15).value, 0.0)

    def test_unavailable_code_is_visible_and_not_converted_to_zero(self):
        model = build_excel_forecast_model(export_payload())
        row = next(item for item in model["sections"][0]["rows"] if item["budget_code"] == "U01")
        self.assertTrue(row["unavailable"])
        self.assertEqual(row["months"], [None] * 12)
        self.assertIsNone(row["ytd"])
        sheet, _ = _sheet(export_payload())
        excel_row = _row_for(sheet, "U01")
        self.assertEqual(sheet.cell(excel_row, 2).value, "Unavailable Code")
        self.assertIsNone(sheet.cell(excel_row, 3).value)
        self.assertIsNone(sheet.cell(excel_row, 15).value)

    def test_unavailable_value_is_excluded_from_category_total(self):
        model = build_excel_forecast_model(export_payload())
        section = next(item for item in model["sections"] if item["category"] == "R&M-AMCs")
        self.assertEqual(section["totals"][0], 10.0 + 30.0 + 0.0)
        self.assertIsNone(next(item for item in section["rows"] if item["budget_code"] == "U01")["months"][0])

    def test_negative_forecast_remains_negative(self):
        model = build_excel_forecast_model(export_payload())
        row = next(item for item in model["sections"][1]["rows"] if item["budget_code"] == "511101")
        self.assertEqual(row["months"][0], -2.5)
        sheet, _ = _sheet(export_payload())
        excel_row = _row_for(sheet, "511101")
        self.assertEqual(sheet.cell(excel_row, 3).value, -2.5)
        self.assertLess(sheet.cell(excel_row, 3).value, 0)

    def test_all_categories_export_creates_separate_sections(self):
        model = build_excel_forecast_model(export_payload())
        self.assertEqual([section["category"] for section in model["sections"]], ["R&M-AMCs", "Int'l Settlement"])
        sheet, _ = _sheet(export_payload())
        self.assertIsNotNone(_row_for(sheet, "R&M-AMCs"))
        self.assertIsNotNone(_row_for(sheet, "Int'l Settlement"))

    def test_single_category_export_contains_only_that_category(self):
        months = _month_range("2027-01", "2027-12")
        payload = export_payload(
            category="R&M-AMCs",
            selected_accounts=["521101", "521102"],
            selected_account_count=2,
            account_monthly_forecasts=(
                _code_rows("521101", "Repair A", months, [10.0] * 12)
                + _code_rows("521102", "Repair B", months, [20.0] * 12)
            ),
            budget_code_forecasts=[
                {
                    "budget_code": "521101",
                    "account_name": "Repair A",
                    "category": "R&M-AMCs",
                    "available": True,
                    "forecast": dict(zip(months, [10.0] * 12)),
                    "production_status": "FITTED_MODEL",
                },
                {
                    "budget_code": "521102",
                    "account_name": "Repair B",
                    "category": "R&M-AMCs",
                    "available": True,
                    "forecast": dict(zip(months, [20.0] * 12)),
                    "production_status": "FITTED_MODEL",
                },
            ],
        )
        model = build_excel_forecast_model(payload)
        self.assertEqual([section["category"] for section in model["sections"]], ["R&M-AMCs"])
        self.assertEqual([row["budget_code"] for row in model["sections"][0]["rows"]], ["521101", "521102"])

    def test_existing_forecast_values_are_unchanged_by_export_generation(self):
        payload = export_payload()
        source = payload["budget_code_forecasts"][0]["forecast"]["2027-01"]
        model = build_excel_forecast_model(payload)
        row = next(item for item in model["sections"][0]["rows"] if item["budget_code"] == "521101")
        self.assertEqual(row["months"][0], source)
        self.assertEqual(row["months"][0], 10.0)
        self.assertEqual(payload["budget_code_forecasts"][0]["forecast"]["2027-01"], 10.0)
        self.assertEqual(payload["account_monthly_forecasts"][0]["forecast_amount"], 10.0)

    def test_excel_filename_is_dynamic(self):
        payload = export_payload()
        self.assertEqual(excel_filename_stem(payload), "AI_Budget_Forecast_Jan_2027_Dec_2027")
        with patch.object(forecasting, "get_forecast_payload", return_value=payload):
            response = forecasting.forecast_report_excel(21)
        self.assertIn('filename="AI_Budget_Forecast_Jan_2027_Dec_2027.xlsx"', response.headers["content-disposition"])

    def test_protected_pickles_were_not_modified(self):
        for path in PROTECTED_PICKLES:
            self.assertEqual(_sha256(path), self.pickle_hashes[str(path)])


if __name__ == "__main__":
    unittest.main()
