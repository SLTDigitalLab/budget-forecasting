"""Focused tests for master-dataset retraining. Training itself is mocked."""

from __future__ import annotations

import os
import pickle
import sys
import tempfile
import time
import unittest
from contextlib import nullcontext
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

BACKEND_DIR = Path(__file__).resolve().parents[2]
ENGINE_DIR = BACKEND_DIR / "model_training_engine"
if str(ENGINE_DIR) not in sys.path:
    sys.path.insert(0, str(ENGINE_DIR))
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import model_training  # noqa: E402
from app.config import FORECAST_ARTIFACT_PATH  # noqa: E402
from app.services import model_loader  # noqa: E402
from app.services.dataset_master import (  # noqa: E402
    MasterSnapshotError,
    resolve_training_snapshot,
)
from app.services.model_loader import (  # noqa: E402
    get_model_bundle,
    load_production_model,
    reset_model_cache,
)

PRODUCTION_STAT = FORECAST_ARTIFACT_PATH.stat()
IS_ARTIFACT = ENGINE_DIR / "output" / "international_settlement_top12_models.pkl"
IS_STAT = IS_ARTIFACT.stat()


def _record(code: str, month: str, amount, description: str, category: str) -> dict:
    return {
        "budget_code": code,
        "month": month,
        "amount": amount,
        "description": description,
        "category": category,
    }


def _database_records() -> list[dict]:
    return [
        _record("511101", "2023-01", 0.0, "Salaries", "Staff Cost"),
        _record("511101", "2023-03", -2.5, "Salaries", "Staff Cost"),
        _record("511102", "2023-01", 10.0, "Fuel", "Vehicle Cost"),
    ]


def _production_bundle(marker: str) -> dict:
    code = "511101"
    return {
        "artifact_schema_version": "all_account_per_budget_code_v1",
        "forecast_type": "PER_BUDGET_CODE_BEST_MODELS",
        "category": "ALL",
        "selected_accounts": [code],
        "overall_best_algorithm": "PER_BUDGET_CODE",
        "historical_start": "Jan-2023",
        "historical_end": marker,
        "evaluation_metrics": ["MASE"],
        "forecast_start": "Jul-2026",
        "forecast_end": "Dec-2027",
        "monthly_account_forecasts": [],
        "monthly_combined_forecast": [],
        "yearly_account_forecasts": [],
        "yearly_combined_forecast": [],
        "year_status": [],
        "account_records": {
            code: {
                "budget_code": code,
                "production_status": "ZERO_POLICY",
                "model": None,
                "trained_model": None,
                "algorithm": None,
            }
        },
        "conformal_prediction": {},
        "forecast_coverage": {"no_data_accounts": [], "zero_policy_accounts": [code]},
        "history_months": ["Jan-2023"],
        "models": {},
    }


def _write_workbook(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Historical Actuals"
    worksheet.append(["ACT CODE", "ACT NAME", "Category", "January23 Actuals", "February23 Actuals"])
    worksheet.append([None, "Staff Cost", "Staff Cost", None, None])
    worksheet.append(["511101", "Salaries", "Staff Cost", 0, -5])
    worksheet.append([None, "Vehicle Cost", "Vehicle Cost", None, None])
    worksheet.append(["511102", "Fuel", "Vehicle Cost", 8, None])
    workbook.save(path)
    workbook.close()


class MasterSnapshotTests(unittest.TestCase):
    def test_workbook_is_used_before_database_initialization(self):
        with tempfile.TemporaryDirectory() as tmp:
            workbook = Path(tmp) / "historical_actuals_master.xlsx"
            _write_workbook(workbook)
            with (
                patch("app.services.dataset_master.list_historical_actuals", return_value=[]),
                patch("app.services.dataset_master.master_version", return_value=0),
                patch(
                    "app.services.dataset_master.load_master_identity",
                    return_value={"budget_codes": [], "earliest_month": None, "latest_month": None},
                ),
                patch("app.services.dataset_master.authoritative_workbook_path", return_value=workbook),
                patch("app.services.dataset_master.export_training_snapshot") as export_snapshot,
            ):
                snapshot = resolve_training_snapshot()
            export_snapshot.assert_not_called()
        self.assertEqual(snapshot["source"], "production_workbook")
        self.assertEqual(snapshot["path"], workbook)
        self.assertEqual(snapshot["budget_code_count"], 2)
        self.assertEqual(snapshot["earliest_month"], "2023-01")
        self.assertEqual(snapshot["latest_month"], "2023-02")
        self.assertNotIn("original_actual_data.xlsx", str(snapshot["path"]))

    def test_database_master_exports_a_training_snapshot(self):
        with tempfile.TemporaryDirectory() as tmp:
            storage = Path(tmp)
            with (
                patch("app.services.dataset_storage.DATASET_STORAGE_DIR", storage),
                patch("app.services.dataset_master.list_historical_actuals", return_value=_database_records()),
                patch("app.services.dataset_master.master_version", return_value=4),
                patch(
                    "app.services.dataset_master.load_master_identity",
                    return_value={"budget_codes": [], "earliest_month": None, "latest_month": None},
                ),
                patch("app.services.dataset_master.authoritative_workbook_path") as workbook_path,
            ):
                snapshot = resolve_training_snapshot()
            workbook_path.assert_not_called()
            self.assertEqual(snapshot["source"], "database")
            self.assertEqual(snapshot["path"], storage / "historical_actuals_master.xlsx")
            self.assertEqual(snapshot["budget_code_count"], 2)
            self.assertEqual(snapshot["earliest_month"], "2023-01")
            self.assertEqual(snapshot["latest_month"], "2023-03")
            saved = load_workbook(snapshot["path"], read_only=True, data_only=True)
            try:
                rows = list(saved.active.iter_rows(values_only=True))
            finally:
                saved.close()
        headers = list(rows[0])
        self.assertIn("February23 Actuals", headers)
        salary = next(row for row in rows if row[0] == "511101")
        february = headers.index("February23 Actuals")
        january = headers.index("January23 Actuals")
        march = headers.index("March23 Actuals")
        self.assertEqual(salary[january], 0)
        self.assertIsNone(salary[february])
        self.assertEqual(salary[march], -2.5)

    def test_failed_database_export_does_not_train_from_another_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            storage = Path(tmp)
            with (
                patch("app.services.dataset_storage.DATASET_STORAGE_DIR", storage),
                patch("app.services.dataset_master.list_historical_actuals", return_value=_database_records()),
                patch("app.services.dataset_master.master_version", return_value=4),
                patch(
                    "app.services.dataset_master.load_master_identity",
                    return_value={"budget_codes": [], "earliest_month": None, "latest_month": None},
                ),
                patch("app.services.dataset_master.publish_managed_master", side_effect=OSError("disk full")),
                patch("app.services.dataset_master.authoritative_workbook_path") as workbook_path,
                patch.object(model_training, "run_training") as run_training,
            ):
                with self.assertRaises(MasterSnapshotError):
                    resolve_training_snapshot()
            workbook_path.assert_not_called()
            run_training.assert_not_called()
            self.assertFalse((storage / "historical_actuals_master.xlsx").exists())


class ProductionPublishTests(unittest.TestCase):
    def test_successful_publication_keeps_the_fixed_filename(self):
        bundle = _production_bundle("Jun-2026")
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            destination = directory / "all_budget_code_models.pkl"
            destination.write_bytes(b"previous-production-artifact")
            temporary = directory / "candidate.pkl.tmp"
            temporary.write_bytes(pickle.dumps(bundle))
            published = model_training.publish_validated_production_pickle(
                temporary,
                destination,
                bundle,
                expected_account_count=1,
            )
            self.assertEqual(published.name, "all_budget_code_models.pkl")
            self.assertEqual(published, destination)
            self.assertFalse(temporary.exists())
            self.assertEqual(pickle.loads(destination.read_bytes())["historical_end"], "Jun-2026")

    def test_invalid_candidate_preserves_the_existing_pickle(self):
        bundle = _production_bundle("Jun-2026")
        del bundle["account_records"]
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            destination = directory / "all_budget_code_models.pkl"
            destination.write_bytes(b"previous-production-artifact")
            temporary = directory / "candidate.pkl.tmp"
            temporary.write_bytes(pickle.dumps(bundle))
            with self.assertRaises(AssertionError):
                model_training.publish_validated_production_pickle(
                    temporary,
                    destination,
                    bundle,
                    expected_account_count=1,
                )
            self.assertEqual(destination.read_bytes(), b"previous-production-artifact")
            self.assertFalse(temporary.exists())

    def test_protected_settlement_artifact_is_not_replaced(self):
        bundle = _production_bundle("Jun-2026")
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            destination = directory / "international_settlement_top12_models.pkl"
            destination.write_bytes(b"settlement-artifact")
            temporary = directory / "candidate.pkl.tmp"
            temporary.write_bytes(pickle.dumps(bundle))
            with self.assertRaises(AssertionError):
                model_training.publish_validated_production_pickle(temporary, destination, bundle)
            self.assertEqual(destination.read_bytes(), b"settlement-artifact")

    def test_failed_training_preserves_the_existing_pickle(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            production = directory / "all_budget_code_models.pkl"
            production.write_bytes(b"previous-production-artifact")
            snapshot = directory / "historical_actuals_master.xlsx"
            snapshot.write_bytes(b"snapshot")
            args = model_training.parse_args(["--retrain-from-master"])

            def fail_training(**_kwargs):
                raise RuntimeError("training failed")

            with (
                patch(
                    "app.services.dataset_master.resolve_training_snapshot",
                    return_value={
                        "source": "production_workbook",
                        "path": snapshot,
                        "budget_code_count": 2,
                        "earliest_month": "2023-01",
                        "latest_month": "2023-02",
                    },
                ),
                patch("app.services.retrain_jobs.cli_retrain_scope", return_value=nullcontext()),
                patch.object(model_training, "run_training", side_effect=fail_training),
            ):
                code = model_training.retrain_from_current_master(args, output_dir=directory)
            self.assertEqual(code, 1)
            self.assertEqual(production.read_bytes(), b"previous-production-artifact")

    def test_retrain_command_uses_the_snapshot_and_fixed_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            snapshot = directory / "historical_actuals_master.xlsx"
            snapshot.write_bytes(b"snapshot")
            args = model_training.parse_args(["--retrain-from-master"])
            captured = {}

            def fake_training(**kwargs):
                captured.update(kwargs)
                path = Path(kwargs["output_dir"]) / kwargs["output_file"]
                path.write_bytes(b"published")
                return {"pickle_verified": True, "pickle_path": path}

            with (
                patch(
                    "app.services.dataset_master.resolve_training_snapshot",
                    return_value={
                        "source": "database",
                        "path": snapshot,
                        "budget_code_count": 2,
                        "earliest_month": "2023-01",
                        "latest_month": "2026-06",
                    },
                ),
                patch("app.services.retrain_jobs.cli_retrain_scope", return_value=nullcontext()),
                patch.object(model_training, "run_training", side_effect=fake_training),
            ):
                code = model_training.retrain_from_current_master(args, output_dir=directory)
            self.assertEqual(code, 0)
            self.assertEqual(captured["input_file"], str(snapshot))
            self.assertNotIn("original_actual_data.xlsx", captured["input_file"])
            self.assertEqual(captured["output_file"], "all_budget_code_models.pkl")
            self.assertEqual(captured["expected_account_count"], 2)
            self.assertTrue((directory / "all_budget_code_models.pkl").is_file())

    def test_snapshot_failure_does_not_start_training(self):
        args = model_training.parse_args(["--retrain-from-master"])
        with tempfile.TemporaryDirectory() as tmp:
            with (
                patch(
                    "app.services.dataset_master.resolve_training_snapshot",
                    side_effect=MasterSnapshotError("export failed"),
                ),
                patch("app.services.retrain_jobs.cli_retrain_scope", return_value=nullcontext()),
                patch.object(model_training, "run_training") as run_training,
            ):
                code = model_training.retrain_from_current_master(args, output_dir=tmp)
            run_training.assert_not_called()
            self.assertEqual(code, 1)

    def test_concurrent_retraining_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            first = model_training.acquire_retrain_lock(directory)
            try:
                with self.assertRaises(model_training.RetrainInProgress):
                    model_training.acquire_retrain_lock(directory)
            finally:
                model_training.release_retrain_lock(*first)
            second = model_training.acquire_retrain_lock(directory)
            model_training.release_retrain_lock(*second)

    def test_stale_lock_can_be_replaced(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            lock = directory / ".all_budget_code_models.retrain.lock"
            lock.write_text("0", encoding="utf-8")
            acquired = model_training.acquire_retrain_lock(directory)
            model_training.release_retrain_lock(*acquired)
            self.assertFalse(lock.exists())


class ForecastReloadTests(unittest.TestCase):
    def setUp(self):
        reset_model_cache()

    def tearDown(self):
        reset_model_cache()

    def test_replaced_artifact_is_loaded_without_dropping_the_previous_object(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "all_budget_code_models.pkl"
            path.write_bytes(pickle.dumps(_production_bundle("Jun-2026")))
            with patch.object(model_loader, "resolve_forecast_artifact_path", return_value=path.resolve()):
                load_production_model()
                previous = get_model_bundle()
                self.assertEqual(previous["historical_end"], "Jun-2026")
                path.write_bytes(pickle.dumps(_production_bundle("Jul-2026-replaced")))
                future = time.time() + 5
                os.utime(path, (future, future))
                updated = get_model_bundle()
            self.assertEqual(previous["historical_end"], "Jun-2026")
            self.assertIsNot(updated, previous)
            self.assertEqual(updated["historical_end"], "Jul-2026-replaced")

    def test_partial_replacement_is_not_loaded(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "all_budget_code_models.pkl"
            path.write_bytes(pickle.dumps(_production_bundle("Jun-2026")))
            with patch.object(model_loader, "resolve_forecast_artifact_path", return_value=path.resolve()):
                load_production_model()
                previous = get_model_bundle()
                path.write_bytes(b"partial-pickle")
                future = time.time() + 5
                os.utime(path, (future, future))
                current = get_model_bundle()
            self.assertIs(current, previous)
            self.assertEqual(current["historical_end"], "Jun-2026")


class ProductionArtifactUntouchedTests(unittest.TestCase):
    def test_real_production_pickles_were_not_modified(self):
        self.assertEqual(FORECAST_ARTIFACT_PATH.stat().st_size, PRODUCTION_STAT.st_size)
        self.assertEqual(FORECAST_ARTIFACT_PATH.stat().st_mtime_ns, PRODUCTION_STAT.st_mtime_ns)
        self.assertEqual(IS_ARTIFACT.stat().st_size, IS_STAT.st_size)
        self.assertEqual(IS_ARTIFACT.stat().st_mtime_ns, IS_STAT.st_mtime_ns)


if __name__ == "__main__":
    unittest.main()
