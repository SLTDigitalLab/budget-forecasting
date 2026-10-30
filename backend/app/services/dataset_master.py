"""Load and atomically refresh the managed historical master workbook."""

from __future__ import annotations

import hashlib
import json
import sys
import uuid
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.config import BACKEND_DIR, DATASET_STORAGE_DIR, HISTORICAL_MASTER_DATASET_PATH, MANAGED_MASTER_FILENAME
from app.services.dataset_parser import actuals_header_for_month, iso_month
from app.services.dataset_repository import list_declarations, list_historical_actuals, load_master_identity, master_version
from app.services.dataset_storage import ensure_storage_dir
from app.services.dataset_workflow import extend_month_span, merge_budget_codes

ENGINE_DIR = BACKEND_DIR / "model_training_engine"


def fingerprint_master_state(state: dict) -> str:
    """Stable identity of the master revision used for a training snapshot."""
    codes = []
    for item in state.get("budget_codes") or []:
        codes.append(
            {
                "budget_code": str(item.get("budget_code")),
                "description": item.get("description") or "",
                "category": item.get("category") or "",
                "effective_from": item.get("effective_from") or None,
                "effective_to": item.get("effective_to") or None,
            }
        )
    records = []
    for item in state.get("records") or []:
        amount = item.get("amount")
        records.append(
            (
                str(item.get("budget_code")),
                str(item.get("month")),
                None if amount is None else float(amount),
            )
        )
    payload = {
        "version": int(state.get("version") or 0),
        "earliest_month": state.get("earliest_month"),
        "latest_month": state.get("latest_month"),
        "budget_codes": sorted(
            [
                {
                    "budget_code": item["budget_code"],
                    "description": item["description"],
                    "category": item["category"],
                }
                for item in codes
            ],
            key=lambda item: item["budget_code"],
        ),
        "records": sorted(records),
    }
    declarations = sorted(
        (str(item.get("budget_code")), str(item.get("month")))
        for item in (state.get("declarations") or [])
        if item.get("budget_code") and item.get("month")
    )
    bounds = sorted(
        (item["budget_code"], item.get("effective_from") or None, item.get("effective_to") or None)
        for item in codes
        if item.get("effective_from") or item.get("effective_to")
    )
    if declarations:
        payload["declarations"] = declarations
    if bounds:
        payload["applicability"] = bounds
    _import_engine()
    from preprocessing import HISTORICAL_ACTUAL_POLICY_VERSION

    payload["historical_actual_policy"] = HISTORICAL_ACTUAL_POLICY_VERSION
    encoded = json.dumps(payload, sort_keys=True, default=str).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def load_master_records() -> tuple[list[dict], int, str]:
    state = load_master_state()
    return state["records"], state["version"], state["source"]


def load_master_state() -> dict:
    """Current master, including budget codes that have no numeric actuals."""
    version = master_version()
    stored = list_historical_actuals()
    identity = load_master_identity()
    if stored or identity["budget_codes"]:
        declarations = list_declarations()
        # Reconcile meta bounds with confirmed actuals/declarations so a stale
        # dataset_master_meta.latest_month cannot hide later confirmed months.
        earliest, latest = extend_month_span(
            identity["earliest_month"],
            identity["latest_month"],
            [*stored, *[{"month": item.get("month")} for item in declarations]],
        )
        codes = merge_budget_codes(identity["budget_codes"], [], stored)
        order = list(identity.get("category_order") or [])
        if not order:
            _import_engine()
            from preprocessing import category_order_from_headings

            order = category_order_from_headings(
                [str(item.get("category") or "") for item in codes]
            )
        return {
            "records": stored,
            "declarations": declarations,
            "budget_codes": codes,
            "category_order": order,
            "earliest_month": earliest,
            "latest_month": latest,
            "version": version,
            "source": "database",
        }
    workbook = master_identity_from_workbook_path(authoritative_workbook_path())
    return {**workbook, "declarations": [], "version": version, "source": "production_workbook"}


def master_identity_from_workbook_path(path: Path) -> dict:
    """Read every budget code and the workbook month range. Missing amounts stay missing."""
    _import_engine()
    from preprocessing import KIND_NUMBER, _collect_accounts, _load_original_excel, assign_categories, parse_amount_cell

    loaded = _load_original_excel(str(path))
    collected = _collect_accounts(loaded)
    assigned = assign_categories(
        loaded["df"],
        loaded["account_code_col"],
        loaded["account_name_col"],
        category_heading_excel_rows=loaded.get("category_heading_excel_rows"),
    )
    accounts = assigned[(~assigned["is_category_row"].fillna(False)) & assigned[loaded["account_code_col"]].notna()]
    metadata = collected["metadata"]
    category_order = list(loaded.get("category_order") or [])
    codes = []
    seen_codes: set[str] = set()
    # Preserve workbook account order for stable category discovery downstream.
    for _, row in accounts.iterrows():
        code = str(row.get("cleaned_account_code") or row.get(loaded["account_code_col"]) or "").strip()
        if not code or code.lower() == "nan" or code in seen_codes:
            continue
        seen_codes.add(code)
        meta = metadata.loc[str(code)] if str(code) in metadata.index else None
        description = "" if meta is None else str(meta.get("account_name") or "")
        category = str(row.get("category") or ("" if meta is None else meta.get("category") or "")).strip()
        if description.lower() == "nan":
            description = ""
        if category.lower() == "nan":
            category = ""
        codes.append({"budget_code": str(code), "description": description, "category": category})
    records = []
    code_column = loaded["account_code_col"]
    for _, row in accounts.iterrows():
        code = str(row.get("cleaned_account_code") or row.get(code_column) or "").strip()
        if not code or code.lower() == "nan":
            continue
        meta = next((item for item in codes if item["budget_code"] == code), None)
        for column, timestamp in zip(loaded["monthly_columns"], loaded["month_dates"]):
            parsed = parse_amount_cell(row.get(column))
            if parsed["kind"] != KIND_NUMBER:
                continue
            records.append(
                {
                    "budget_code": code,
                    "month": iso_month(timestamp),
                    "amount": float(parsed["value"]),
                    "description": "" if meta is None else meta["description"],
                    "category": "" if meta is None else meta["category"],
                    "source_file_id": None,
                }
            )
    months = [iso_month(timestamp) for timestamp in loaded["month_dates"]]
    if not codes or not months:
        raise MasterSnapshotError(f"The historical workbook has no budget codes: {path}")
    return {
        "records": records,
        "declarations": [],
        "budget_codes": codes,
        "category_order": category_order,
        "earliest_month": months[0],
        "latest_month": months[-1],
    }


def records_from_production_workbook() -> list[dict]:
    import sys

    from app.config import BACKEND_DIR

    engine = BACKEND_DIR / "model_training_engine"
    if str(engine) not in sys.path:
        sys.path.insert(0, str(engine))
    from preprocessing import KIND_NUMBER, parse_amount_cell

    from app.routers.overview import _load_workbook

    workbook = _load_workbook()
    assigned = workbook["assigned"]
    code_column = workbook["account_code_col"]
    name_column = workbook["account_name_col"]
    records = []
    accounts = assigned.loc[~assigned["is_category_row"].fillna(False) & assigned["numeric_account_code"].notna()]
    for _, row in accounts.iterrows():
        code = str(row.get("cleaned_account_code") or "").strip()
        if not code or code.lower() == "nan":
            continue
        description = str(row.get(name_column) or "").strip()
        category = str(row.get("category") or "").strip()
        if category.lower() == "nan":
            category = ""
        for column, timestamp in zip(workbook["monthly_columns"], workbook["month_dates"]):
            parsed = parse_amount_cell(row.get(column))
            if parsed["kind"] != KIND_NUMBER:
                continue
            records.append(
                {
                    "budget_code": code,
                    "month": iso_month(timestamp),
                    "amount": float(parsed["value"]),
                    "description": description,
                    "category": category,
                    "source_file_id": None,
                }
            )
    if not records:
        raise ValueError("The production historical workbook did not contain usable monthly actuals.")
    return records


class MasterSnapshotError(RuntimeError):
    """The current historical master could not be prepared for training."""


def write_managed_master(records: list[dict]) -> Path:
    temp = prepare_managed_master(records)
    return publish_managed_master(temp)


def export_training_snapshot(
    records: list[dict],
    *,
    budget_codes: list[dict] | None = None,
    category_order: list[str] | None = None,
    earliest_month: str | None = None,
    latest_month: str | None = None,
    declarations: list[dict] | None = None,
) -> Path:
    """Write the database master to the managed workbook and prove training can read it."""
    try:
        temp = prepare_managed_master(
            records,
            budget_codes=budget_codes,
            category_order=category_order,
            earliest_month=earliest_month,
            latest_month=latest_month,
            declarations=declarations,
        )
    except Exception as error:
        raise MasterSnapshotError(
            "Database historical master could not be exported. Retraining was not started."
        ) from error
    try:
        ensure_training_snapshot(temp)
        return publish_managed_master(temp)
    except Exception as error:
        temp.unlink(missing_ok=True)
        raise MasterSnapshotError(
            "Database historical master could not be exported. Retraining was not started."
        ) from error


def prepare_managed_master(
    records: list[dict],
    *,
    budget_codes: list[dict] | None = None,
    category_order: list[str] | None = None,
    earliest_month: str | None = None,
    latest_month: str | None = None,
    declarations: list[dict] | None = None,
) -> Path:
    grouped: dict[str, dict] = {}
    months: set[str] = set()
    for item in budget_codes or []:
        code = str(item.get("budget_code") or "").strip()
        if not code:
            continue
        grouped.setdefault(
            code,
            {
                "description": item.get("description") or "",
                "category": item.get("category") or "",
                "amounts": {},
            },
        )
    if earliest_month:
        months.add(str(earliest_month))
    if latest_month:
        months.add(str(latest_month))
    for item in records:
        if item.get("amount") is None:
            continue
        code = str(item["budget_code"])
        bucket = grouped.setdefault(
            code,
            {
                "description": item.get("description") or "",
                "category": item.get("category") or "",
                "amounts": {},
            },
        )
        if item.get("description"):
            bucket["description"] = item["description"]
        if item.get("category"):
            bucket["category"] = item["category"]
        month = str(item["month"])
        months.add(month)
        bucket["amounts"][month] = float(item["amount"])
    for item in declarations or []:
        code = str(item.get("budget_code") or "").strip()
        month = str(item.get("month") or "").strip()
        if not code or not month:
            continue
        bucket = grouped.setdefault(
            code,
            {"description": "", "category": "", "amounts": {}},
        )
        if month not in bucket["amounts"]:
            bucket["amounts"][month] = "-"
            months.add(month)
    if not grouped or not months:
        raise ValueError("The historical master has no numeric actuals to export.")
    ordered_months = _training_months(sorted(months))
    directory = ensure_storage_dir()
    temp = directory / f".tmp-master-{uuid.uuid4().hex}.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Historical Actuals"
    headers = ["ACT CODE", "ACT NAME", "Category", *[actuals_header_for_month(month) for month in ordered_months]]
    worksheet.append(headers)
    from openpyxl.styles import Font

    _import_engine()
    from preprocessing import category_order_from_headings, sort_categories_by_order

    by_category: dict[str, list[str]] = {}
    for code in sorted(grouped):
        by_category.setdefault(str(grouped[code]["category"] or ""), []).append(code)
    discovered = category_order_from_headings([name for name in by_category if name])
    ordered_categories = sort_categories_by_order(discovered, category_order)
    # Keep uncategorized accounts last, without inventing a heading.
    if "" in by_category:
        ordered_categories = [*ordered_categories, ""]
    heading_font = Font(bold=True)
    category_headers = 0
    for category in ordered_categories:
        if category:
            category_headers += 1
            # Heading format: blank ACT CODE, bold ACT NAME, empty Category + actuals cells.
            worksheet.append([None, category, None, *([None] * len(ordered_months))])
            name_cell = worksheet.cell(worksheet.max_row, 2)
            name_cell.font = heading_font
        for code in by_category.get(category, []):
            bucket = grouped[code]
            worksheet.append(
                [
                    code,
                    bucket["description"],
                    bucket["category"] or None,
                    *[bucket["amounts"].get(month) for month in ordered_months],
                ]
            )
    workbook.save(temp)
    workbook.close()
    _verify_workbook(temp, len(grouped), len(ordered_months), category_headers)
    return temp


def publish_managed_master(temp: Path) -> Path:
    directory = ensure_storage_dir()
    target = directory / MANAGED_MASTER_FILENAME
    backup = directory / f"{MANAGED_MASTER_FILENAME}.bak"
    if target.exists():
        backup.write_bytes(target.read_bytes())
    temp.replace(target)
    return target


def _verify_workbook(path: Path, code_count: int, month_count: int, category_header_count: int = 0) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    column_count = len(rows[0]) if rows else 0
    if len(rows) != code_count + category_header_count + 1 or column_count != month_count + 3:
        path.unlink(missing_ok=True)
        raise ValueError("The managed historical master workbook failed verification and was not replaced.")


def managed_master_path() -> Path:
    return Path(DATASET_STORAGE_DIR) / MANAGED_MASTER_FILENAME


def scan_workbook_explicit_hyphens(path: Path) -> dict:
    """Read stored cell values. A displayed accounting dash is not a declaration."""
    from openpyxl import load_workbook

    from app.services.dataset_parser import iso_month

    _import_engine()
    from preprocessing import (  # noqa: WPS433
        clean_account_code,
        is_excluded_actuals_column,
        normalize_label,
        parse_month_header,
    )

    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        header = None
        header_index = None
        for index, row in enumerate(worksheet.iter_rows()):
            if any(isinstance(cell.value, str) and "ACT CODE" in cell.value.upper() for cell in row):
                header = row
                header_index = index
                break
        if header is None:
            return {"hyphen_pairs": [], "numeric_zero_dash_format": 0}
        headers = [normalize_label(cell.value) for cell in header]
        code_index = next(index for index, label in enumerate(headers) if label.upper() == "ACT CODE" or "ACT CODE" in label.upper())
        month_indexes: list[tuple[int, str]] = []
        for index, label in enumerate(headers):
            if not label or is_excluded_actuals_column(label):
                continue
            timestamp = parse_month_header(label)
            if pd_notna(timestamp):
                month_indexes.append((index, iso_month(timestamp)))
        pairs: list[tuple[str, str]] = []
        zero_dash = 0
        for row in worksheet.iter_rows(min_row=header_index + 2):
            code = clean_account_code(row[code_index].value)
            if not isinstance(code, str) or not code.strip():
                continue
            for index, month in month_indexes:
                cell = row[index]
                value = cell.value
                if isinstance(value, str) and value.strip() == "-":
                    pairs.append((str(code), month))
                elif value == 0 or value == 0.0:
                    if "-" in str(cell.number_format or ""):
                        zero_dash += 1
        return {"hyphen_pairs": pairs, "numeric_zero_dash_format": zero_dash}
    finally:
        workbook.close()


def pd_notna(value) -> bool:
    try:
        import pandas as pd
    except ImportError:
        return value == value
    return bool(pd.notna(value))


def extract_original_baseline(path: Path) -> dict:
    """Read the approved original workbook. Blanks in supplied months become declarations.

    This does not apply to uploads. The source file is not modified.
    """
    import hashlib

    from openpyxl import load_workbook

    from app.services.dataset_parser import iso_month
    from app.services.submission_state import SUBMISSION_DECLARED_MISSING, SUBMISSION_NUMERIC, classify_baseline_stored_cell

    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    _import_engine()
    from preprocessing import (
        KIND_NUMBER,
        _load_original_excel,
        assign_categories,
        clean_account_code,
        normalize_label,
        parse_amount_cell,
    )

    loaded = _load_original_excel(str(path))
    assigned = assign_categories(
        loaded["df"],
        loaded["account_code_col"],
        loaded["account_name_col"],
        category_heading_excel_rows=loaded.get("category_heading_excel_rows"),
    )
    header_row = loaded["header_row_idx"] + 1

    def _sheet_rows(data_only: bool) -> tuple[list, list]:
        book = load_workbook(path, data_only=data_only, read_only=True)
        try:
            sheet = book[loaded["sheet_name"]]
            rows = list(sheet.iter_rows(min_row=header_row))
        finally:
            book.close()
        if not rows:
            return [], []
        return rows[0], rows[1:]

    stored_header, stored_rows = _sheet_rows(False)
    cached_header, cached_rows = _sheet_rows(True)
    stored_headers = [normalize_label(cell.value) for cell in stored_header]
    column_index = {label: index for index, label in enumerate(stored_headers)}
    month_positions = []
    for column, timestamp in zip(loaded["monthly_columns"], loaded["month_dates"]):
        label = normalize_label(column)
        if label not in column_index:
            continue
        month_positions.append((column_index[label], iso_month(timestamp)))
    del cached_header
    skip_labels = {"total", "grand total", "subtotal"}
    records = []
    declarations = []
    ambiguous = []
    codes = []
    seen = set()
    numeric_pairs: set[tuple[str, str]] = set()
    code_column = loaded["account_code_col"]
    name_column = loaded["account_name_col"]
    index_by_month = {month: index for index, month in month_positions}
    for _, row in assigned.iterrows():
        if bool(row.get("is_category_row")):
            continue
        code = clean_account_code(row.get(code_column))
        if not isinstance(code, str) or not code.strip() or code.casefold() in skip_labels:
            continue
        description = "" if row.get(name_column) is None else str(row.get(name_column)).strip()
        if description.lower() == "nan":
            description = ""
        category = "" if row.get("category") is None else str(row.get("category")).strip()
        if category.lower() == "nan":
            category = ""
        if description.casefold() in skip_labels:
            continue
        source_row = int(row["_source_row"])
        offset = source_row - (loaded["header_row_idx"] + 2)
        if code not in seen:
            seen.add(code)
            codes.append(
                {
                    "budget_code": code,
                    "description": description,
                    "category": category,
                    "effective_from": None,
                    "effective_to": None,
                }
            )
        for column, timestamp in zip(loaded["monthly_columns"], loaded["month_dates"]):
            month = iso_month(timestamp)
            parsed = parse_amount_cell(row.get(column))
            if parsed["kind"] == KIND_NUMBER:
                amount = float(parsed["value"])
                records.append(
                    {
                        "budget_code": code,
                        "month": month,
                        "amount": amount,
                        "description": description,
                        "category": category,
                    }
                )
                numeric_pairs.add((code, month))
                continue
            index = index_by_month.get(month)
            stored_value = None
            cached_value = None
            if index is None or offset < 0 or offset >= len(stored_rows):
                ambiguous.append(
                    {
                        "budget_code": code,
                        "month": month,
                        "reason": "supplied month column was not readable",
                    }
                )
                continue
            stored_row = stored_rows[offset]
            cached_row = cached_rows[offset] if offset < len(cached_rows) else ()
            stored_value = stored_row[index].value if index < len(stored_row) else None
            cached_value = cached_row[index].value if index < len(cached_row) else None
            classified = classify_baseline_stored_cell(stored_value, cached_value)
            if classified["kind"] == SUBMISSION_NUMERIC and (code, month) not in numeric_pairs:
                records.append(
                    {
                        "budget_code": code,
                        "month": month,
                        "amount": float(classified["value"]),
                        "description": description,
                        "category": category,
                    }
                )
                numeric_pairs.add((code, month))
            elif classified["kind"] == SUBMISSION_DECLARED_MISSING and (code, month) not in numeric_pairs:
                declarations.append({"budget_code": code, "month": month})
            elif (code, month) not in numeric_pairs:
                ambiguous.append(
                    {
                        "budget_code": code,
                        "month": month,
                        "reason": classified.get("reason") or "ambiguous baseline cell",
                    }
                )
    months = [iso_month(timestamp) for timestamp in loaded["month_dates"]]
    return {
        "baseline_hash": digest,
        "records": records,
        "declarations": declarations,
        "budget_codes": codes,
        "ambiguous": ambiguous,
        "earliest_month": months[0] if months else None,
        "latest_month": months[-1] if months else None,
        "numeric_count": len(records),
        "declaration_count": len(declarations),
        "budget_code_count": len(codes),
    }


def authoritative_workbook_path() -> Path:
    """Workbook used before historical actuals are stored in the database."""
    dataset_dir = ENGINE_DIR / "dataset"
    candidates = [
        Path(HISTORICAL_MASTER_DATASET_PATH),
        dataset_dir / "historical_actuals_master.xlsx",
        dataset_dir / "original_actual_data.xlsx",
    ]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise MasterSnapshotError("The authoritative historical workbook was not found.")


def resolve_training_snapshot() -> dict:
    """Return the current master snapshot. A failed database export does not use an older file."""
    state = load_master_state()
    if state["source"] == "database":
        path = export_training_snapshot(
            state["records"],
            budget_codes=state["budget_codes"],
            earliest_month=state["earliest_month"],
            latest_month=state["latest_month"],
            declarations=state.get("declarations") or [],
        )
        summary = {
            "budget_code_count": len(state["budget_codes"]),
            "earliest_month": state["earliest_month"],
            "latest_month": state["latest_month"],
        }
        source = "database"
        version = state["version"]
    else:
        path = authoritative_workbook_path()
        summary = _summary_from_workbook(path)
        source = "production_workbook"
        version = state["version"]
    if not path.is_file() or path.stat().st_size <= 0:
        raise MasterSnapshotError(f"Historical master snapshot is missing: {path}")
    return {
        "source": source,
        "path": path,
        "master_version": version,
        **summary,
    }


def _summary_from_records(records: list[dict]) -> dict:
    codes = {str(item["budget_code"]) for item in records if item.get("amount") is not None and item.get("budget_code")}
    months = sorted({str(item["month"]) for item in records if item.get("amount") is not None and item.get("month")})
    if not codes or not months:
        raise MasterSnapshotError("The database historical master has no numeric actuals.")
    return {
        "budget_code_count": len(codes),
        "earliest_month": months[0],
        "latest_month": months[-1],
    }


def _summary_from_workbook(path: Path) -> dict:
    _import_engine()
    from preprocessing import _collect_accounts, _load_original_excel

    loaded = _load_original_excel(str(path))
    collected = _collect_accounts(loaded)
    months = list(loaded["month_dates"])
    if collected["account_count"] <= 0 or not months:
        raise MasterSnapshotError(f"The historical workbook has no trainable accounts: {path}")
    start = months[0]
    end = months[-1]
    return {
        "budget_code_count": int(collected["account_count"]),
        "earliest_month": f"{int(start.year):04d}-{int(start.month):02d}",
        "latest_month": f"{int(end.year):04d}-{int(end.month):02d}",
    }


def _training_months(observed_months: list[str]) -> list[str]:
    _import_engine()
    from preprocessing import HISTORY_START

    start_key = f"{int(HISTORY_START.year):04d}-{int(HISTORY_START.month):02d}"
    early = [month for month in observed_months if month < start_key]
    if early:
        raise ValueError(
            "The historical master contains months before the training calendar and was not exported."
        )
    latest = observed_months[-1]
    if latest < start_key:
        raise ValueError("The historical master ends before the training calendar and was not exported.")
    year, month = int(HISTORY_START.year), int(HISTORY_START.month)
    end_year, end_month = (int(part) for part in latest.split("-"))
    columns: list[str] = []
    while (year, month) <= (end_year, end_month):
        columns.append(f"{year:04d}-{month:02d}")
        month += 1
        if month == 13:
            month = 1
            year += 1
    return columns


def ensure_training_snapshot(path: Path) -> None:
    """Fail when the exported workbook is not readable by the training parser."""
    _import_engine()
    from preprocessing import _load_original_excel

    _load_original_excel(str(path))


def _import_engine() -> None:
    if str(ENGINE_DIR) not in sys.path:
        sys.path.insert(0, str(ENGINE_DIR))
