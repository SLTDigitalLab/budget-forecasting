"""Excel workbook for a persisted forecast report."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.worksheet import Worksheet

from app.services.report_builder import UNAVAILABLE, build_excel_forecast_model

HEADER_FILL = PatternFill("solid", fgColor="071A33")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="071A33")
CATEGORY_FONT = Font(bold=True, size=12, color="071A33")
TOTAL_FONT = Font(bold=True, color="071A33")
CATEGORY_FILL = PatternFill("solid", fgColor="D6E3F0")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
SECTION_FILL = PatternFill("solid", fgColor="023E8A")
SECTION_FONT = Font(bold=True, color="FFFFFF", size=12)
AMOUNT_FORMAT = "#,##0.00;(#,##0.00)"
SHEET_NAME = "Budget Forecast"
SHEET_NAMES = (SHEET_NAME,)
# Coverage description uses ACT CODE + ACT NAME widths; count sits in the next column.
COVERAGE_LABEL_END_COL = 2
COVERAGE_COUNT_COL = 3


def _set_forecast_widths(sheet: Worksheet, month_count: int) -> None:
    """Set main forecast table widths only — do not stretch month columns for summaries."""
    widths = [18, 44, *[12] * month_count, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_numeric_or_blank(cell, value: Any) -> None:
    if value == UNAVAILABLE or value is None:
        cell.value = None
        cell.alignment = Alignment(horizontal="right", vertical="center")
        return
    cell.value = float(value)
    cell.number_format = AMOUNT_FORMAT
    cell.alignment = Alignment(horizontal="right", vertical="center")


def _month_end_col(month_count: int) -> int:
    return 2 + month_count


def _ytd_col(month_count: int) -> int:
    return 3 + month_count


def _sum_formula(start_col: int, start_row: int, end_col: int, end_row: int) -> str:
    start = f"{get_column_letter(start_col)}{start_row}"
    end = f"{get_column_letter(end_col)}{end_row}"
    return f"=SUM({start}:{end})"


def _ensure_editable(workbook: Workbook, sheet: Worksheet) -> None:
    """Export real editable cells — no sheet/workbook locks or read-only recommendation."""
    # Do not assign protection.password = None; openpyxl hashes the value and raises.
    sheet.protection.sheet = False
    if getattr(sheet.protection, "password", None):
        sheet.protection.set_password("", already_hashed=True)
    workbook.security = WorkbookProtection(
        workbookPassword=None,
        revisionsPassword=None,
        lockStructure=False,
        lockWindows=False,
        lockRevision=False,
    )


def render_forecast_excel(report: dict[str, Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    _write_budget_forecast(sheet, report)
    _ensure_editable(workbook, sheet)
    extra = [name for name in workbook.sheetnames if name != SHEET_NAME]
    if extra:
        raise AssertionError("Excel export must contain only the Budget Forecast worksheet.")
    if workbook.sheetnames != [SHEET_NAME]:
        raise AssertionError("Excel workbook must contain exactly one worksheet: Budget Forecast.")
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_budget_forecast(sheet: Worksheet, report: dict[str, Any]) -> None:
    sheet.title = SHEET_NAME
    model = report.get("excel_forecast") or build_excel_forecast_model({
        "requested_start_month": report.get("requested_start_month"),
        "requested_end_month": report.get("requested_end_month"),
        "category": (report.get("header") or {}).get("category"),
        "selected_accounts": report.get("selected_accounts") or [],
        "monthly_forecasts": report.get("monthly_forecasts") or report.get("chart_months") or [],
        "account_monthly_forecasts": report.get("account_monthly_forecasts") or [],
        "budget_code_forecasts": report.get("budget_code_forecasts") or [],
    })
    months = list(model.get("months") or [])
    month_count = len(months)
    headers = ["ACT CODE", "ACT NAME", *list(model.get("month_headers") or []), "YTD"]
    last_col = len(headers)
    month_end = _month_end_col(month_count)
    ytd_col = _ytd_col(month_count)

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(last_col, 1))
    title = sheet.cell(1, 1, model.get("title") or "AI-Based Budget Forecast")
    title.font = TITLE_FONT
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 26

    sections = list(model.get("sections") or [])
    partial = any(
        item.get("unavailable")
        for section in sections
        for item in (section.get("rows") or [])
    )
    if partial:
        note = sheet.cell(
            2,
            1,
            "Forecast unavailable codes have blank amounts. Those blanks are not zeros. "
            "Totals exclude them and are partial.",
        )
        note.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(last_col, 1))
        sheet.row_dimensions[2].height = 36

    row = 3
    for section in sections:
        category = str(section.get("category") or "")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        heading = sheet.cell(row, 1, category)
        heading.font = CATEGORY_FONT
        heading.fill = CATEGORY_FILL
        heading.alignment = Alignment(vertical="center", wrap_text=True)
        for column in range(1, last_col + 1):
            sheet.cell(row, column).fill = CATEGORY_FILL
            sheet.cell(row, column).border = THIN
        sheet.row_dimensions[row].height = 22
        row += 1

        for column, label in enumerate(headers, start=1):
            cell = sheet.cell(row, column, label)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN
        sheet.row_dimensions[row].height = 20
        row += 1

        section_partial = any(item.get("unavailable") for item in (section.get("rows") or []))
        data_rows = list(section.get("rows") or [])
        data_start = row
        for item in data_rows:
            code_cell = sheet.cell(row, 1, str(item.get("budget_code") or ""))
            code_cell.number_format = "@"
            code_cell.alignment = Alignment(horizontal="left", vertical="center")
            account_name = str(item.get("account_name") or "")
            if item.get("unavailable"):
                account_name = f"{account_name} — Forecast unavailable".strip()
            name_cell = sheet.cell(row, 2, account_name)
            name_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            values = list(item.get("months") or [])
            for index in range(month_count):
                amount = values[index] if index < len(values) else None
                _write_numeric_or_blank(sheet.cell(row, 3 + index), amount)
            if month_count:
                ytd_cell = sheet.cell(row, ytd_col)
                ytd_cell.value = _sum_formula(3, row, month_end, row)
                ytd_cell.number_format = AMOUNT_FORMAT
                ytd_cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                sheet.cell(row, ytd_col).value = None
            for column in range(1, last_col + 1):
                sheet.cell(row, column).border = THIN
            sheet.row_dimensions[row].height = 20
            row += 1
        data_end = row - 1

        total_label = sheet.cell(row, 1, "TOTAL (partial)" if section_partial else "TOTAL")
        total_label.font = TOTAL_FONT
        sheet.cell(row, 2, "")
        if data_rows and month_count:
            for index in range(month_count):
                total_cell = sheet.cell(row, 3 + index)
                total_cell.value = _sum_formula(3 + index, data_start, 3 + index, data_end)
                total_cell.number_format = AMOUNT_FORMAT
                total_cell.alignment = Alignment(horizontal="right", vertical="center")
            ytd_total = sheet.cell(row, ytd_col)
            ytd_total.value = _sum_formula(ytd_col, data_start, ytd_col, data_end)
            ytd_total.number_format = AMOUNT_FORMAT
            ytd_total.alignment = Alignment(horizontal="right", vertical="center")
        else:
            for index in range(month_count):
                _write_numeric_or_blank(sheet.cell(row, 3 + index), None)
            sheet.cell(row, ytd_col).value = None
        for column in range(1, last_col + 1):
            sheet.cell(row, column).fill = TOTAL_FILL
            sheet.cell(row, column).border = THIN
            sheet.cell(row, column).font = TOTAL_FONT
        sheet.row_dimensions[row].height = 20
        row += 2

    if not sections:
        message = sheet.cell(row, 1, model.get("no_forecast_message") or "No forecast available for the selected category.")
        message.font = Font(bold=True, color="071A33")
        message.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(last_col, 1))
        sheet.row_dimensions[row].height = 28
        row += 2

    row = _write_coverage_summary(sheet, model, row, last_col)
    row += 1
    _write_unavailable_accounts(sheet, model, row, last_col)
    _set_forecast_widths(sheet, month_count)
    sheet.freeze_panes = "A3"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _section_banner(sheet: Worksheet, row: int, last_col: int, title: str) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(last_col, 1))
    cell = sheet.cell(row, 1, title)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for column in range(1, max(last_col, 1) + 1):
        sheet.cell(row, column).fill = SECTION_FILL
        sheet.cell(row, column).font = SECTION_FONT
        sheet.cell(row, column).border = THIN
    sheet.row_dimensions[row].height = 24


def _coverage_row_height(label: str) -> float:
    text = str(label or "")
    # ~55 visible chars across ACT CODE + ACT NAME widths; grow for wrap.
    lines = max(1, (len(text) // 55) + 1)
    return max(28, min(48, 16 + lines * 14))


def _write_coverage_summary(sheet: Worksheet, model: dict[str, Any], row: int, last_col: int) -> int:
    _section_banner(sheet, row, last_col, "FORECAST COVERAGE SUMMARY")
    row += 1

    label_end = min(COVERAGE_LABEL_END_COL, max(last_col, 1))
    count_col = min(COVERAGE_COUNT_COL, max(last_col, 1))
    if count_col <= label_end:
        count_col = label_end + 1 if last_col > label_end else label_end

    if label_end > 1:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=label_end)
    description_header = sheet.cell(row, 1, "Description")
    description_header.font = HEADER_FONT
    description_header.fill = HEADER_FILL
    description_header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column in range(1, label_end + 1):
        sheet.cell(row, column).font = HEADER_FONT
        sheet.cell(row, column).fill = HEADER_FILL
        sheet.cell(row, column).border = THIN
        sheet.cell(row, column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    count_header = sheet.cell(row, count_col, "Budget Codes")
    count_header.font = HEADER_FONT
    count_header.fill = HEADER_FILL
    count_header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    count_header.border = THIN
    sheet.row_dimensions[row].height = 22
    row += 1

    for label, count in model.get("coverage_rows") or []:
        if label_end > 1:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=label_end)
        description = sheet.cell(row, 1, str(label))
        description.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for column in range(1, label_end + 1):
            sheet.cell(row, column).border = THIN
            sheet.cell(row, column).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        value = sheet.cell(row, count_col, int(count))
        value.alignment = Alignment(horizontal="right", vertical="center")
        value.number_format = "0"
        value.border = THIN
        sheet.row_dimensions[row].height = _coverage_row_height(str(label))
        row += 1
    return row


def _unavailable_row_height(reason: str, account_name: str) -> float:
    longest = max(len(str(reason or "")), len(str(account_name or "")), 1)
    lines = max(1, (longest // 40) + 1)
    return max(24, min(60, 14 + lines * 12))


def _write_unavailable_accounts(sheet: Worksheet, model: dict[str, Any], row: int, last_col: int) -> None:
    _section_banner(sheet, row, last_col, "EXCLUDED & UNAVAILABLE BUDGET CODES")
    row += 1
    accounts = list(model.get("unavailable_accounts") or [])
    if not accounts:
        note = sheet.cell(row, 1, model.get("no_unavailable_message") or "No excluded or unavailable budget codes.")
        note.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(last_col, 1))
        sheet.row_dimensions[row].height = 28
        return

    headers = ["ACT CODE", "ACT NAME", "CATEGORY", "STATUS", "REASON"]
    reason_col = 5
    reason_end = max(last_col, reason_col)
    for column, label in enumerate(headers, start=1):
        if column == reason_col and reason_end > reason_col:
            sheet.merge_cells(start_row=row, start_column=reason_col, end_row=row, end_column=reason_end)
        cell = sheet.cell(row, column, label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    for column in range(reason_col, reason_end + 1):
        sheet.cell(row, column).font = HEADER_FONT
        sheet.cell(row, column).fill = HEADER_FILL
        sheet.cell(row, column).border = THIN
    sheet.row_dimensions[row].height = 22
    row += 1

    for item in accounts:
        code_cell = sheet.cell(row, 1, str(item.get("budget_code") or ""))
        code_cell.number_format = "@"
        code_cell.alignment = Alignment(horizontal="left", vertical="center")
        account_name = str(item.get("account_name") or "")
        sheet.cell(row, 2, account_name).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )
        sheet.cell(row, 3, str(item.get("category") or "")).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )
        sheet.cell(row, 4, "UNAVAILABLE").alignment = Alignment(horizontal="left", vertical="center")
        reason = str(item.get("reason") or "")
        if reason_end > reason_col:
            sheet.merge_cells(start_row=row, start_column=reason_col, end_row=row, end_column=reason_end)
        sheet.cell(row, reason_col, reason).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )
        for column in range(1, reason_end + 1):
            sheet.cell(row, column).border = THIN
        sheet.row_dimensions[row].height = _unavailable_row_height(reason, account_name)
        row += 1
