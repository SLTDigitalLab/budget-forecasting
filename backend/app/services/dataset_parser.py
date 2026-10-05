"""Wide-format historical dataset parsing and validation."""

from __future__ import annotations

import sys
import uuid
from collections import defaultdict
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from app.config import BACKEND_DIR

ENGINE_DIR = BACKEND_DIR / "model_training_engine"
if str(ENGINE_DIR) not in sys.path:
    sys.path.insert(0, str(ENGINE_DIR))

from preprocessing import (  # noqa: E402
    KIND_INVALID,
    KIND_MISSING,
    KIND_NUMBER,
    assign_categories,
    clean_account_code,
    is_excluded_actuals_column,
    label_key,
    normalize_label,
    parse_amount_cell,
    parse_month_header,
)

CODE_ALIASES = ("ACT CODE", "BUDGET CODE", "ACCOUNT CODE")
NAME_ALIASES = ("ACT NAME", "DESCRIPTION", "ACCOUNT NAME")
CATEGORY_ALIASES = ("CATEGORY",)
MAX_HEADER_SCAN_ROWS = 100
MAX_SHEETS = 40
MAX_ROWS = 20000
MAX_COLUMNS = 250


class DatasetParseError(Exception):
    pass


def iso_month(timestamp: pd.Timestamp) -> str:
    value = pd.Timestamp(timestamp).replace(day=1)
    return f"{int(value.year):04d}-{int(value.month):02d}"


def actuals_header_for_month(month: str) -> str:
    year, month_number = str(month).split("-")
    name = pd.Timestamp(year=int(year), month=int(month_number), day=1).strftime("%B")
    return f"{name}{year[2:]} Actuals"


def _lookup_column(columns: list[str], aliases: tuple[str, ...]) -> str | None:
    lookup = {label_key(column): column for column in columns}
    for alias in aliases:
        if alias in lookup:
            return lookup[alias]
    return None


def detect_upload_month_columns(column_names: list[str]) -> tuple[list[str], list[pd.Timestamp], list[dict]]:
    detected: list[tuple[str, pd.Timestamp]] = []
    seen: dict[pd.Timestamp, str] = {}
    errors: list[dict] = []
    for column in column_names:
        label = normalize_label(column)
        if not label:
            continue
        excluded = is_excluded_actuals_column(label)
        timestamp = parse_month_header(label)
        looks_monthly = "actual" in label.lower() or bool(pd.notna(timestamp))
        if excluded:
            continue
        if looks_monthly and pd.isna(timestamp):
            errors.append(
                {
                    "sheet": "",
                    "row": 1,
                    "column": label,
                    "message": f"Monthly header {label!r} is invalid or ambiguous.",
                    "budget_code": "",
                }
            )
            continue
        if pd.isna(timestamp):
            continue
        month = pd.Timestamp(timestamp).replace(day=1)
        if month in seen:
            errors.append(
                {
                    "sheet": "",
                    "row": 1,
                    "column": label,
                    "message": (
                        f"Multiple columns resolve to {iso_month(month)}: "
                        f"{seen[month]!r} and {label!r}."
                    ),
                    "budget_code": "",
                }
            )
            continue
        seen[month] = label
        detected.append((label, month))
    detected.sort(key=lambda item: item[1])
    return [item[0] for item in detected], [item[1] for item in detected], errors


def _excel_row_number(header_row_idx: int, data_index: int) -> int:
    return header_row_idx + 2 + int(data_index)


def _is_blank(value: Any) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return True
    return str(value).strip() == ""


def _workbook_formula_errors(path: Path, sheet_name: str, header_row_idx: int, columns: list[str]) -> list[dict]:
    errors: list[dict] = []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            return errors
        worksheet = workbook[sheet_name]
        header = [normalize_label(cell.value) for cell in next(worksheet.iter_rows(min_row=header_row_idx + 1, max_row=header_row_idx + 1))]
        col_index = {label: idx for idx, label in enumerate(header) if label}
        for excel_row, row in enumerate(
            worksheet.iter_rows(min_row=header_row_idx + 2, max_row=header_row_idx + 1 + MAX_ROWS),
            start=header_row_idx + 2,
        ):
            for column in columns:
                idx = col_index.get(column)
                if idx is None or idx >= len(row):
                    continue
                cell = row[idx]
                if cell.value is None:
                    continue
                if getattr(cell, "data_type", "") == "f" or (
                    isinstance(cell.value, str) and str(cell.value).startswith("=")
                ):
                    errors.append(
                        {
                            "sheet": sheet_name,
                            "row": excel_row,
                            "column": column,
                            "message": (
                                f"Cell {column} on row {excel_row} contains a formula or unsupported workbook feature. "
                                "Replace it with a numeric actual before saving."
                            ),
                            "budget_code": "",
                        }
                    )
    finally:
        workbook.close()
    return errors


def list_candidate_sheets(frame_by_sheet: dict[str, pd.DataFrame]) -> list[dict]:
    candidates = []
    for sheet_name, preview in frame_by_sheet.items():
        for idx, row in preview.iterrows():
            labels = [label_key(value) for value in row.tolist()]
            if any(alias in labels for alias in CODE_ALIASES):
                candidates.append({"sheet": sheet_name, "header_row": int(idx) + 1})
                break
    return candidates


def _read_tabular(path: Path, filename: str, sheet_name: str | None) -> dict[str, Any]:
    suffix = path.suffix.lower()
    original_name = filename.lower()
    if suffix in {".csv"} or original_name.endswith(".csv"):
        raw = pd.read_csv(path, header=None, dtype=object, nrows=MAX_ROWS + MAX_HEADER_SCAN_ROWS)
        return {
            "kind": "csv",
            "sheets": ["Sheet1"],
            "selected_sheet": "Sheet1",
            "raw": raw,
            "header_row_idx": 0,
            "needs_sheet_selection": False,
            "candidates": [{"sheet": "Sheet1", "header_row": 1}],
        }
    if suffix not in {".xlsx", ".xlsm"} and not original_name.endswith((".xlsx", ".xlsm")):
        raise DatasetParseError("Only .xlsx and .csv files are supported.")

    xls = pd.ExcelFile(path)
    sheets = list(xls.sheet_names)[:MAX_SHEETS]
    previews = {}
    candidates = []
    for candidate in sheets:
        preview = pd.read_excel(xls, sheet_name=candidate, header=None, dtype=object, nrows=MAX_HEADER_SCAN_ROWS)
        previews[candidate] = preview
        found = False
        for idx, row in preview.iterrows():
            labels = [label_key(value) for value in row.tolist()]
            if any(alias in labels for alias in CODE_ALIASES):
                candidates.append({"sheet": candidate, "header_row": int(idx) + 1})
                found = True
                break
        if not found:
            continue
    if not candidates:
        raise DatasetParseError("No worksheet contained a Budget Code / ACT CODE header in the first 100 rows.")
    if sheet_name:
        match = next((item for item in candidates if item["sheet"] == sheet_name), None)
        if match is None:
            raise DatasetParseError(f"Sheet {sheet_name!r} is not a valid historical dataset sheet.")
        selected = match["sheet"]
    elif len(candidates) > 1:
        return {
            "kind": "xlsx",
            "sheets": sheets,
            "candidates": candidates,
            "needs_sheet_selection": True,
            "selected_sheet": None,
            "raw": None,
            "header_row_idx": None,
        }
    else:
        selected = candidates[0]["sheet"]

    raw = pd.read_excel(path, sheet_name=selected, header=None, dtype=object, nrows=MAX_ROWS + MAX_HEADER_SCAN_ROWS)
    header_row_idx = next(item["header_row"] for item in candidates if item["sheet"] == selected) - 1
    return {
        "kind": "xlsx",
        "sheets": sheets,
        "candidates": candidates,
        "needs_sheet_selection": False,
        "selected_sheet": selected,
        "raw": raw,
        "header_row_idx": header_row_idx,
    }


def parse_historical_dataset(path: Path, filename: str, sheet_name: str | None = None) -> dict[str, Any]:
    loaded = _read_tabular(path, filename, sheet_name)
    if loaded.get("needs_sheet_selection"):
        return {
            "ok": False,
            "needs_sheet_selection": True,
            "sheets": loaded["candidates"],
            "errors": [
                {
                    "sheet": "",
                    "row": None,
                    "column": "",
                    "message": "Multiple sheets look like historical datasets. Select one sheet to validate.",
                    "budget_code": "",
                }
            ],
            "rows": [],
            "months": [],
            "duplicate_codes": [],
            "can_confirm": False,
            "selected_sheet": None,
            "filename": filename,
        }

    raw = loaded["raw"]
    header_row_idx = int(loaded["header_row_idx"])
    sheet = loaded["selected_sheet"]
    headers = [normalize_label(value) for value in raw.iloc[header_row_idx].tolist()]
    if len(headers) > MAX_COLUMNS:
        headers = headers[:MAX_COLUMNS]
    df = raw.iloc[header_row_idx + 1 : header_row_idx + 1 + MAX_ROWS].copy()
    df.columns = headers
    df = df.loc[:, [column for column in df.columns if str(column).strip() != ""]]
    df = df.reset_index(drop=True)

    errors: list[dict] = []
    code_col = _lookup_column(df.columns.tolist(), CODE_ALIASES)
    name_col = _lookup_column(df.columns.tolist(), NAME_ALIASES)
    category_col = _lookup_column(df.columns.tolist(), CATEGORY_ALIASES)
    if not code_col:
        errors.append(
            {
                "sheet": sheet,
                "row": header_row_idx + 1,
                "column": "Budget Code",
                "message": "Required column Budget Code / ACT CODE is missing.",
                "budget_code": "",
            }
        )
    if not name_col:
        errors.append(
            {
                "sheet": sheet,
                "row": header_row_idx + 1,
                "column": "Description",
                "message": "Required column Description / ACT NAME is missing.",
                "budget_code": "",
            }
        )
    month_columns, month_dates, month_errors = detect_upload_month_columns(df.columns.tolist())
    for item in month_errors:
        item["sheet"] = sheet
        errors.append(item)
    if not month_columns and not month_errors:
        errors.append(
            {
                "sheet": sheet,
                "row": header_row_idx + 1,
                "column": "",
                "message": "No monthly actual columns were detected. Use headers such as July25 Actuals.",
                "budget_code": "",
            }
        )
    if loaded["kind"] == "xlsx":
        inspect_columns = [column for column in [code_col, name_col, category_col, *month_columns] if column]
        errors.extend(_workbook_formula_errors(path, sheet, header_row_idx, inspect_columns))

    if errors and (not code_col or not name_col or not month_columns):
        return {
            "ok": False,
            "needs_sheet_selection": False,
            "sheets": loaded["candidates"],
            "errors": errors,
            "rows": [],
            "months": [iso_month(item) for item in month_dates],
            "duplicate_codes": [],
            "can_confirm": False,
            "selected_sheet": sheet,
            "filename": filename,
        }

    working = df.copy()
    working[code_col] = working[code_col].map(clean_account_code)
    working[name_col] = working[name_col].map(
        lambda value: "" if _is_blank(value) else str(value).strip()
    )
    if category_col:
        working["category"] = working[category_col].map(
            lambda value: "" if _is_blank(value) else str(value).strip()
        )
        working["is_category_row"] = False
    else:
        assigned = assign_categories(working, code_col, name_col)
        working = assigned
        working["category"] = working["category"].fillna("").astype(str).str.strip()

    parsed_rows: list[dict] = []
    codes_to_rows: dict[str, list[int]] = defaultdict(list)
    for index, record in working.iterrows():
        excel_row = _excel_row_number(header_row_idx, int(index))
        if record.get("is_category_row"):
            continue
        code = record.get(code_col)
        description = str(record.get(name_col) or "").strip()
        category = str(record.get("category") or "").strip()
        amounts: dict[str, Any] = {}
        raw_amounts: dict[str, Any] = {}
        missing_months: list[str] = []
        invalid_months: list[str] = []
        invalid_amounts = []
        all_amounts_blank = True
        for column, timestamp in zip(month_columns, month_dates):
            raw_value = record.get(column)
            month = iso_month(timestamp)
            raw_amounts[month] = None if _is_blank(raw_value) else raw_value
            parsed_cell = parse_amount_cell(raw_value)
            if parsed_cell["kind"] == KIND_MISSING:
                amounts[month] = None
                missing_months.append(month)
                continue
            all_amounts_blank = False
            if parsed_cell["kind"] == KIND_INVALID:
                invalid_months.append(month)
                invalid_amounts.append((excel_row, column, month, raw_value))
                amounts[month] = None
            else:
                amounts[month] = float(parsed_cell["value"])
        blank_row = _is_blank(code) and not description and not category and all_amounts_blank
        if blank_row:
            continue
        if _is_blank(code):
            errors.append(
                {
                    "sheet": sheet,
                    "row": excel_row,
                    "column": code_col,
                    "message": f"Row {excel_row} is missing a Budget Code.",
                    "budget_code": "",
                }
            )
            continue
        code_text = str(code).strip()
        codes_to_rows[code_text].append(excel_row)
        if not description:
            errors.append(
                {
                    "sheet": sheet,
                    "row": excel_row,
                    "column": name_col,
                    "message": f"Budget code {code_text} is missing a description on row {excel_row}.",
                    "budget_code": code_text,
                }
            )
        if not category:
            errors.append(
                {
                    "sheet": sheet,
                    "row": excel_row,
                    "column": category_col or "Category",
                    "message": f"Budget code {code_text} is missing a category on row {excel_row}.",
                    "budget_code": code_text,
                }
            )
        for row_number, column, month, raw_value in invalid_amounts:
            errors.append(
                {
                    "sheet": sheet,
                    "row": row_number,
                    "column": column,
                    "message": (
                        f"Budget code {code_text} has a non-numeric amount for {month} "
                        f"in column {column} on row {row_number}."
                    ),
                    "budget_code": code_text,
                }
            )
        parsed_rows.append(
            {
                "stage_id": uuid.uuid4().hex,
                "budget_code": code_text,
                "description": description,
                "category": category,
                "source_row": excel_row,
                "amounts": amounts,
                "raw_amounts": raw_amounts,
                "missing_months": missing_months,
                "invalid_months": invalid_months,
            }
        )

    month_hits: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in parsed_rows:
        for month, amount in (row.get("amounts") or {}).items():
            if amount is None or month in (row.get("invalid_months") or []):
                continue
            month_hits[(row["budget_code"], month)].append(
                {
                    "stage_id": row["stage_id"],
                    "source_row": row["source_row"],
                    "amount": amount,
                    "description": row.get("description") or "",
                    "category": row.get("category") or "",
                }
            )
    duplicate_conflicts = []
    for (code, month), hits in month_hits.items():
        if len(hits) < 2:
            continue
        duplicate_conflicts.append({"budget_code": code, "month": month, "records": hits})
        errors.append(
            {
                "sheet": sheet,
                "row": hits[0]["source_row"],
                "column": code_col,
                "message": (
                    f"Budget code {code} appears more than once for {month} "
                    f"(rows {', '.join(str(item['source_row']) for item in hits)}). "
                    "Correct the duplicates before confirming. Rows were not combined."
                ),
                "budget_code": code,
            }
        )
    duplicate_codes = [
        {"budget_code": code, "rows": rows}
        for code, rows in codes_to_rows.items()
        if len(rows) > 1 and any(item["budget_code"] == code for item in duplicate_conflicts)
    ]

    blocking = [item for item in errors if item.get("message")]
    return {
        "ok": not blocking,
        "needs_sheet_selection": False,
        "sheets": loaded["candidates"],
        "errors": errors,
        "rows": parsed_rows,
        "months": [iso_month(item) for item in month_dates],
        "duplicate_codes": duplicate_codes,
        "duplicate_conflicts": duplicate_conflicts,
        "can_confirm": not blocking and bool(parsed_rows),
        "selected_sheet": sheet,
        "filename": filename,
        "header_row_idx": header_row_idx,
        "code_column": code_col,
        "name_column": name_col,
        "category_column": category_col,
        "month_columns": month_columns,
    }
